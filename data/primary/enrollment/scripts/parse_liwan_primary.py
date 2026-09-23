#!/usr/bin/env python3
"""复现荔湾小学层转录 parsed/_transcripts/liwan_2026.json。

数据源（raw/ 官方原文件，2026-04-28 挂网）：
  liwan_2026_primary_a1.docx  附件1 2026年荔湾区公办小学招生计划（序号/学校/班数）
  liwan_2026_primary_a4.docx  附件4 2026年荔湾区公办小学适龄儿童入学登记服务地段划分表（学校/街道/社区/路街巷）

输出结构（与海珠/天河转录一致，消费方兼容）：
  {district, year, source, note, schools: [{school, plan_classes, zone}]}

zone 格式：每行 "街道·社区：地址"（学校列空=上一所学校地段的延续行，街道列空则继承上一条），与旧 B 层产物格式一致。
用法: python3 data/primary/transition/scripts/parse_liwan_primary.py
"""
import json
import os
import re

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(HERE))))
RAW = os.path.join(ROOT, "data", "primary", "enrollment", "raw")
OUT = os.path.join(ROOT, "data", "primary", "enrollment", "parsed", "_transcripts", "liwan_2026.json")

A1 = os.path.join(RAW, "liwan_2026_primary_a1.docx")
A4 = os.path.join(RAW, "liwan_2026_primary_a4.docx")


def clean_cell(v):
    """官方 docx 单元格含换行拆段（如「广州市西关外国语学校附属\n流花小学」）与
    「（注：…）」入学备注（属 zone/备注，不应混入校名）→ 去空白拼接 + 截断（注。"""
    v = re.sub(r"\s+", "", v or "")
    return v.split("（注")[0].strip()


def table_rows(path):
    d = Document(path)
    assert len(d.tables) == 1, f"{path} 表格数 != 1"
    return [[c.text.strip() for c in row.cells] for row in d.tables[0].rows]


def build():
    # 附件1 招生计划
    plan = {}
    for row in table_rows(A1):
        if len(row) >= 3 and row[1] and row[1] != "学校":
            m = re.match(r"^\d+$", row[0])
            if m:
                cls = re.sub(r"\s+", "", str(row[2])) if row[2] else ""
                plan[clean_cell(row[1])] = int(cls) if cls.isdigit() else None
    # 附件4 服务地段划分表
    zones = {}
    cur = None
    last_street = ""
    for row in table_rows(A4):
        if len(row) < 4:
            continue
        school, street, comm, addr = clean_cell(row[0]), clean_cell(row[1]), clean_cell(row[2]), clean_cell(row[3])
        if school and school != "学校":
            cur = school
        if street:
            last_street = street
        if not cur:
            continue
        line = f"{last_street}·{comm}：{addr}" if comm else f"{last_street}：{addr}"
        zones.setdefault(cur, []).append(line)
    schools = []
    for school, zl in zones.items():
        schools.append({
            "school": school,
            "plan_classes": plan.get(school),
            "zone": "\n".join(zl),
        })
    out = {
        "district": "荔湾区", "year": 2026,
        "source": "荔湾区教育局《2026年广州市荔湾区公办小学一年级招生工作方案》附件1（招生计划）+附件4（适龄儿童入学登记服务地段划分表）",
        "note": "python-docx 直读 raw 原文件转录（2026-09-21），zone 每行 街道·社区：地址",
        "schools": schools,
    }
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(f"✓ liwan_2026.json（{len(schools)} 校）→ {OUT}")
    return out


if __name__ == "__main__":
    build()
