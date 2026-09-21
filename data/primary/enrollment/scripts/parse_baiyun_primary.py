#!/usr/bin/env python3
"""复现白云小学层转录 parsed/_transcripts/baiyun_2026.json。

数据源（raw/ 官方原文件）：baiyun_2026_official.xlsx，sheet「公办小学」
  列: 序号/片/街（镇）/学校名称/学校类别（公或民办）/招生服务地段/招生咨询电话/2026年计划招生/特殊情况说明
  只收录 类别==公办 的行（民办行排除）。

输出结构（与各区转录一致）：
  {district, year, source, note, schools: [{school, plan_classes, zone, phone, note}]}

用法: python3 data/primary/enrollment/scripts/parse_baiyun_primary.py
"""
import json
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(HERE))))
RAW = os.path.join(ROOT, "data", "primary", "enrollment", "raw", "baiyun_2026_official.xlsx")
OUT = os.path.join(ROOT, "data", "primary", "enrollment", "parsed", "_transcripts", "baiyun_2026.json")


def build():
    wb = openpyxl.load_workbook(RAW, read_only=True)
    assert "公办小学" in wb.sheetnames, wb.sheetnames
    ws = wb["公办小学"]
    schools = []
    for row in ws.iter_rows(values_only=True):
        vals = [str(c).strip() if c is not None else "" for c in row]
        if len(vals) < 9 or vals[3] == "学校名称":
            continue
        if vals[4] != "公办":
            continue
        schools.append({
            "school": vals[3].replace("\n", "").replace("\r", ""),
            "plan_classes": int(vals[7]) if vals[7].isdigit() else None,
            "zone": vals[5],
            "phone": vals[6],
            "note": vals[8],
        })
    out = {
        "district": "白云区", "year": 2026,
        "source": "白云区教育局《广州市白云区2026年义务教育阶段学校招生计划》附表1：公办小学（含小区配套学校）",
        "note": "openpyxl 直读 raw 原文件 xlsx 转录（2026-09-21），只收类别=公办",
        "schools": schools,
    }
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(f"✓ baiyun_2026.json（{len(schools)} 校）→ {OUT}")
    return out


if __name__ == "__main__":
    build()
