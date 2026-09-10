#!/usr/bin/env python3
"""构建白云区公办小学招生离线数据库 2026。

数据源:
  白云区教育局《广州市白云区2026年义务教育阶段学校招生计划》附件 xlsx
    https://www.by.gov.cn/zwfw/zdfw/xwsq/zcwj/content/post_10791741.html
    附件: https://www.by.gov.cn/attachment/8/8015/8015798/10791741.xlsx
  sheet「公办小学」列: 序号/片/街镇/学校名称/类别/招生服务地段/电话/2026计划招生/特殊说明
  只收录 类别==公办 的行（民办行排除）。

用法: python3 scripts/primary/build_baiyun_2026.py [xlsx路径]
输出: data/primary/enrollments/_raw/baiyun_2026.json  +  2026-baiyun.json
"""
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DATA = os.path.join(ROOT, "data", "primary")
OUT_DIR = os.path.join(DATA, "enrollments")
RAW_DIR = os.path.join(OUT_DIR, "_raw")

# 复用五区通用匹配规则
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_district_enrollment import (  # noqa: E402
    norm_school, rank_candidates,
)

ADCODE = "440111"
DISTRICT = "白云区"

SOURCE = ("广州市白云区教育局《广州市白云区2026年义务教育阶段学校招生计划》"
          "附表1：2026年白云区义务教育学校招生计划表（公办小学含小区配套学校）")
SOURCE_URL = "https://www.by.gov.cn/zwfw/zdfw/xwsq/zcwj/content/post_10791741.html"

# 白云区专属名称映射（官方名 → POI 名）：通用规则匹配不到的别名/新校/跨校区
NAME_MAP = {
    "广州市白云区大岡小学": "大冈小学",
    "广州市白云区横沙小学": "横沙学校",
    "广州市白云区夏良小学": "夏良学校",
    "广州市白云区白云外国语小学": "广州市白云区白云外国语中小学",
    "羊城铁路总公司广州铁路第五小学": "广州铁路第五小学",
    "羊城铁路总公司广州铁路第八小学": "",  # POI 无对应校，置空走 unmatched
    "广州市白云区民航学校\n（校本部）": "白云区民航学校小学部",
    "广州市白云区民航学校（校本部）": "白云区民航学校小学部",
    "广州市白云区民航学校\n（人和校区）": "",  # 同一 POI 已被校本部占用，不重复
    "广州市白云区人和镇第六小学": "人和第六小学",
    "广州市白云区太和第二小学": "太和第二小学(米龙校区)",
    "广州市白云区良田第三小学": "良田第三小学(白沙校区)",
    "广州市白云区竹料第二小学": "竹料第二小学(竹三校区)",
    "广州市白云区金广实验学校\n(金域蓝湾校区）": "金广实验小学",
    "广州市白云区金广实验学校（御金沙校区）": "",  # 同 POI 已占用
    "广州市白云区龙归学校": "龙归学校(小学部)",
    "广州市白云区龙归学校\n（珑璟校区）": "",  # 同 POI 已占用
    "广州市白云区张村中心小学": "石井张村中心小学",
    "广州市白云区新楼小学": "新楼村梁庆贤学校",
    "广州市白云中学\n（棠景校区）": "",  # 白云中学棠景校区（初中附设小学），勿配白云中学附属小学
    "广州市白云区穗丰学校": "广州市白云区穗丰学校",
    "广州市白云区人和镇第一小学（鸦湖校区）": "人和镇第一小学(鸦湖校区)",
    "广州市白云区人和镇第二小学（建南校区）": "广州市白云区人和镇第二小学(建南校区)",
    "广州市白云区六中实验小学\n（南校区）": "白云区六中实验小学(南校区)",
    "广州市白云区六中实验小学\n（北校区）": "白云六中实验小学",
    "广州市白云区人和镇第五小学\n（凤和校区）": "人和镇第五小学",
    "广州市白云区人和镇第五小学\n（太成校区）": "人和镇第五小学(太成校区)",
    "广州市白云区广州空港实验小学": "广州空港实验小学(总校区)",
    "广州市白云区棠溪小学": "白云区棠溪小学(棠溪校区)",
    "广州市白云区汇侨第一小学": "汇侨第一小学(汇侨校区)",
    "广州市白云区景云小学": "",
    "广州市白云区明德小学": "明德小学(明德校区)",
    "广州市白云区金沙小学": "金沙小学(校本部)",
    "广州市白云区沙凤小学": "沙凤小学",
    "广州市白云区三元里小学": "三元里小学",
    "广州市白云区景泰小学": "景泰小学(景泰校区)",
    "广州市白云区远景小学": "远景小学",
    "广州市白云区黄边小学": "黄边小学(黄边校区)",
    "广州市白云区京溪小学": "京溪小学(京溪校区)",
    "广州市白云区华师附中实验小学": "华师附中实验小学(北校区)",
    "广州市白云区潭岗小学": "潭岗小学(潭村校区)",
    "广州市白云区黄石学校": "黄石学校(黄石校区)",
    "广州市白云区东平学校": "东平学校(东平校区)",
    "广州市白云区第六十五中学附属小学": "广州市第六十五中学附属小学",
    "广州市第六十五中学附属小学": "广州市第六十五中学附属小学",
    "广州市白云区培英实验学校\n（小学部）": "广州培英实验学校小学部",
    "广州市第七十三中学小学部": "广州市第七十三中学·小学部",
    "广州市白云区龙岗学校": "龙岗学校-龙岗小学",
    "广州市白云区沙龙小学": "神山镇沙龙小学",
}


def clean(v):
    if v is None:
        return ""
    return re.sub(r"[ \t]*\n[ \t]*", "\n", str(v)).strip()


def extract_records(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["公办小学"]
    records = []
    last_pian, last_street = "", ""
    for r in range(4, ws.max_row + 1):
        no = ws.cell(r, 1).value
        pian = ws.cell(r, 2).value
        street = ws.cell(r, 3).value
        name = ws.cell(r, 4).value
        nature = ws.cell(r, 5).value
        zone = ws.cell(r, 6).value
        phone = ws.cell(r, 7).value
        plan = ws.cell(r, 8).value
        note = ws.cell(r, 9).value
        if pian:
            last_pian = str(pian).strip()
        if street:
            last_street = str(street).strip()
        if not name:
            continue
        name = clean(name)
        nature = str(nature or "").strip()
        if nature != "公办":
            continue
        # plan_classes
        if isinstance(plan, (int, float)):
            plan_classes = int(plan)
        elif isinstance(plan, str) and plan.strip().isdigit():
            plan_classes = int(plan.strip())
        else:
            plan_classes = None
        z = clean(zone)
        n = clean(note)
        rec = {
            "school": name,
            "district": DISTRICT,
            "nature": "公办",
            "plan_classes": plan_classes,
            "zone": z,
            "note": n,
            "source": "白云区教育局2026",
        }
        records.append(rec)
    return records


def match_and_write(records):
    with open(os.path.join(DATA, "schools-gz.json"), encoding="utf-8") as f:
        data = json.load(f)
    poi_pool = [s for s in data.get("schools", []) if s.get("adcode") == ADCODE]
    BAD_POI = ("建设中", "在建", "工地", "装修", "筹备", "规划", "选址", "暂停营业", "鲸go")
    poi_pool = [s for s in poi_pool if len(s["name"]) > 2 and not any(b in s["name"] for b in BAD_POI)]

    bindings = []
    map_fail = []
    for idx, rec in enumerate(records):
        mapped = NAME_MAP.get(rec["school"])
        if mapped is None:
            cands = rank_candidates(rec["school"], poi_pool)
            if cands:
                bindings.append((cands[0][0], idx, cands[0][2]))
            else:
                map_fail.append(rec["school"])
            continue
        if mapped == "":
            # 显式标记：该官方校不分配 POI（同 POI 已被主校区占用 / 无 POI）
            continue
        poi = next((p for p in poi_pool if p["name"] == mapped), None)
        if poi:
            bindings.append((1100, idx, poi))
        else:
            map_fail.append(rec["school"])

    bindings.sort(key=lambda x: -x[0])
    matched, ambiguous, used = [], [], {}
    for score, idx, poi in bindings:
        if poi["name"] in used:
            ambiguous.append({"school": records[idx]["school"], "poi": poi["name"],
                              "compete_with": used[poi["name"]]})
            continue
        used[poi["name"]] = records[idx]["school"]
        rec = dict(records[idx])
        rec["school_id"] = poi["name"]
        rec["lng"], rec["lat"] = poi["lng"], poi["lat"]
        matched.append(rec)

    poi_leftover = [s["name"] for s in poi_pool if s["name"] not in used]
    matched_names = {r["school"] for r in matched}
    unmatched = [r["school"] for r in records if r["school"] not in matched_names]

    result = {
        "year": 2026, "district": DISTRICT,
        "source": SOURCE, "source_url": SOURCE_URL,
        "records": matched, "unmatched": sorted(set(unmatched)),
        "ambiguous": ambiguous, "poi_leftover": sorted(set(poi_leftover)),
        "map_fail": sorted(set(map_fail)),
    }
    out = os.path.join(OUT_DIR, "2026-baiyun.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=1)

    print(f"[白云区] 官方公办记录 {len(records)} | 匹配 {len(matched)} / 未匹配 {len(unmatched)} / 歧义 {len(ambiguous)} / POI剩余 {len(poi_leftover)}")
    print("  未匹配:")
    for u in sorted(set(unmatched)):
        print("   -", repr(u))
    if ambiguous:
        print("  歧义:", ambiguous)
    return result


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else "/tmp/baiyun2026/plan.xlsx"
    records = extract_records(xlsx)
    # 落 _raw 中间数据
    os.makedirs(RAW_DIR, exist_ok=True)
    with open(os.path.join(RAW_DIR, "baiyun_2026.json"), "w", encoding="utf-8") as f:
        json.dump({"schools": records}, f, ensure_ascii=False, indent=1)
    match_and_write(records)


if __name__ == "__main__":
    main()
